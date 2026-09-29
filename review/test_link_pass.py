# -*- coding: utf-8 -*-
"""test_link_pass.py -- can the citation linker damage a slide, and do its
guards actually fire?

    python3 review/test_link_pass.py

Linking a path that sits inside a sentence means taking one run apart and
rebuilding it from clones. That is the only operation in this pipeline that
can silently change what a slide says, so the text-preservation invariant is
tested against a deliberately broken splitter: if a wrong splitter does not
fail the test, the invariant is not testing anything.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import link_pass as lp


def buggy_split(para, run, spans, mutate):
    """split_and_link with `mutate` applied to the segment list.

    Used only to prove the text-preservation check can fail. It mirrors the
    real splitter closely enough that a mutation here is the same shape as a
    mistake there.
    """
    from copy import deepcopy
    from pptx.oxml.ns import qn
    text = run.text
    segs, pos = [], 0
    for start, end, href in spans:
        if start > pos:
            segs.append((text[pos:start], None))
        segs.append((text[start:end], href))
        pos = end
    if pos < len(text):
        segs.append((text[pos:], None))
    segs = mutate(segs)

    r = run._r
    parent = r.getparent()
    idx = list(parent).index(r)
    fresh = []
    for seg_text, _ in segs:
        nr = deepcopy(r)
        nr.find(qn("a:t")).text = seg_text
        fresh.append(nr)
    parent.remove(r)
    for off, nr in enumerate(fresh):
        parent.insert(idx + off, nr)


def fresh_paragraph(text):
    d = Presentation()
    sl = d.slides.add_slide(d.slide_layouts[6])
    tb = sl.shapes.add_textbox(Pt(10), Pt(10), Pt(400), Pt(100))
    para = tb.text_frame.paragraphs[0]
    run = para.add_run()
    run.text = text
    return para, run


SENTENCES = [
    u"Simulated in ngspice from sim/buck.cir and scripts/headtohead.py.",
    u"sim/dpt.cir at the start of the line.",
    u"At the end of the line, models/egan.lib",
    u"Three: sim/buck.cir, models/segdrv.lib, rtl/seg_gate_ctrl.v together.",
    u"The switch node peaks at 124 V/ns, which is not a path.",
    u"No citation here at all.",
]

# (text, must `broken` flag something?)
BROKEN_CASES = [
    (u"see scripts/does_not_exist.py for the sweep", True),
    (u"see sim/buck.cir for the converter", False),
    (u"peaking at 124 V/ns", False),
    (u"github.com/Amritha902/gan-driver", False),
    (u"Fig. 14 - Simulated in ngspice.", False),
    (u"100 V / 10 A / 25 C", False),
]


def main():
    links = lp.Links()
    bad = 0

    print("  -- text is preserved exactly when a run is split --")
    for text in SENTENCES:
        para, run = fresh_paragraph(text)
        spans = links.spans(run.text)
        if spans:
            lp.split_and_link(para, run, spans)
        after = "".join(r.text for r in para.runs)
        linked = [r.text for r in para.runs if r.hyperlink.address]
        ok = after == text
        bad += 0 if ok else 1
        print("     %-4s %-2d linked  %s" % ("PASS" if ok else "FAIL",
                                             len(linked), text[:58]))

    print("  -- a wrong splitter must be caught, or the check is decorative --")
    # Passing fewer spans is NOT a bug: the tail is still emitted, so the
    # text survives and the invariant correctly stays quiet. The bug class
    # worth testing is a splitter that forgets the trailing segment, which
    # is how a caption would lose its last sentence.
    for label, mutate in [
            ("drops the text after the last citation", lambda segs: segs[:-1]),
            ("drops the text before the first", lambda segs: segs[1:]),
            ("doubles a segment", lambda segs: segs + segs[-1:])]:
        text = SENTENCES[0]
        para, run = fresh_paragraph(text)
        before = "".join(r.text for r in para.runs)
        buggy_split(para, run, links.spans(run.text), mutate)
        after = "".join(r.text for r in para.runs)
        caught = after != before
        bad += 0 if caught else 1
        print("     %-4s %s" % ("PASS" if caught else "FAIL", label))

    print("  -- a cited file that git does not track must fail the build --")
    for text, should in BROKEN_CASES:
        hits = links.broken(text)
        ok = bool(hits) == should
        bad += 0 if ok else 1
        print("     %-4s %-26s %s" % ("PASS" if ok else "FAIL",
                                      hits[0] if hits else "clean", text[:48]))

    print("  -- a real file resolves, an invented one does not --")
    for cand, should in [("sim/buck.cir", True), ("models/egan.lib", True),
                         ("rtl/seg_gate_ctrl.v", True),
                         ("sim/nope.cir", False)]:
        got = [h for _, _, h in links.spans("see %s here" % cand)]
        ok = bool(got) == should
        bad += 0 if ok else 1
        print("     %-4s %-24s %s" % ("PASS" if ok else "FAIL", cand,
                                      got[0] if got else "not linked"))

    print("\n  %s" % ("all checks behaved as expected" if not bad
                      else "%d CHECK(S) FAILED" % bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
