# -*- coding: utf-8 -*-
"""prune_orphans.py -- drop the slides the running order dropped.

rebuild_pass.py picks the deck out of a larger one: ORDER names what ships and
everything else is "dropped". It drops them from the slide id list, which is
what PowerPoint reads to decide what to show -- but the relationship from the
presentation part to each dropped slide stays behind, so the slide itself, its
text, its figures and its media are all still inside the .pptx.

Nothing displays them. That was the reason to leave it alone, and it was a bad
reason: a deliverable that carries 45 retired slides carries their numbers too,
including ones the deck was corrected to stop claiming. Anyone who opens the
file with a repair tool, converts it, or unzips it can read them, and the
RETIRED list in check_consistency.py exists precisely because those numbers
must not travel with the deck.

    python3 review/prune_orphans.py deck.pptx [deck.pptx ...]

Dropping the relationship makes the slide part unreachable from the package
root, and python-pptx serialises only what it can reach -- so the slide, and
any media or image only that slide used, are gone from the saved file. Parts a
surviving slide still needs are reachable through it and stay.
"""
import os
import sys

from pptx import Presentation


def _parts(path):
    """What the saved file actually contains, by part name."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        return {i.filename: i.file_size for i in z.infolist()}


def prune(path, verbose=True):
    """Drop every slide relationship the slide id list does not name."""
    prs = Presentation(path)
    before_parts = _parts(path)
    before_size = os.path.getsize(path)
    shown = len(prs.slides)
    titles = [_title(s) for s in prs.slides]

    keep = {sid.rId for sid in prs.slides._sldIdLst}
    orphans = [rId for rId, rel in prs.part.rels.items()
               if rel.reltype.endswith("/slide") and rId not in keep]
    for rId in orphans:
        prs.part.drop_rel(rId)
    prs.save(path)

    after = Presentation(path)
    if len(after.slides) != shown or [_title(s) for s in after.slides] != titles:
        raise SystemExit("%s: pruning changed the running order -- refusing to "
                         "call that a clean-up" % os.path.basename(path))

    after_parts = _parts(path)
    gone = sorted(set(before_parts) - set(after_parts))
    if verbose:
        print("  %-42s %2d shown, %d orphan slide(s) dropped"
              % (os.path.basename(path), shown, len(orphans)))
        print("     parts %d -> %d,  %.2f MB -> %.2f MB"
              % (len(before_parts), len(after_parts),
                 before_size / 1e6, os.path.getsize(path) / 1e6))
        media = [g for g in gone if "/media/" in g]
        if media:
            print("     media released: %s"
                  % ", ".join(g.split("/")[-1] for g in media))
    return len(orphans)


def _title(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ""


def main(argv):
    if not argv:
        raise SystemExit(__doc__.strip().splitlines()[0])
    for path in argv:
        prune(path)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
